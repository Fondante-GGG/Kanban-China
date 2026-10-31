from backend.application.finance_periods import finance_transaction, require_open_date, business_date
import json
import io
from datetime import date
from decimal import Decimal, InvalidOperation

from django.db.models import Sum, Count, Q
from django.db import transaction
from django.http import JsonResponse, FileResponse, HttpResponse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt

from backend.domain.exceptions import PermissionDeniedException, ValidationException, EntityNotFoundException
from backend.infrastructure.persistence.models import DepartmentModel, ExpenseModel, ExpenseAttachmentModel, ReimbursementPaymentModel, ExpenseActivityLogModel, NotificationModel, UserModel
from backend.presentation.middlewares import login_required_domain


FINANCE_ROLES = {"general_director", "accountant"}
EDITABLE_STATUSES = {"draft", "needs_info"}
VALID_STATUSES = {"draft", "under_review", "needs_info", "approved", "rejected", "paid"}
ALLOWED_STATUS_TRANSITIONS = {
    "draft": {"under_review"},
    "needs_info": {"under_review"},
    "under_review": {"needs_info", "approved", "rejected"},
    "approved": set(),
    "rejected": set(),
    "paid": set(),
}
TERMINAL_ATTACHMENT_STATUSES = {"approved", "rejected", "paid"}


def _can_manage(user):
    return str(user.role) in FINANCE_ROLES


def _can_access(user, expense):
    if _can_manage(user) or expense.created_by_id == user.id:
        return True
    return str(user.role) == "manager" and user.department_id and expense.department_id == user.department_id


def _serialize(expense, detailed=False):
    data = {
        "id": expense.id, "purchase_date": expense.purchase_date.isoformat(), "department_id": expense.department_id,
        "department_name": expense.department.name if expense.department else None, "category": expense.category,
        "supplier": expense.supplier, "description": expense.description, "amount": str(expense.amount),
        "currency": expense.currency, "payment_source": expense.payment_source, "payment_method": expense.payment_method,
        "status": expense.status, "created_by_id": expense.created_by_id, "created_by_name": expense.created_by.full_name,
        "accountant_comment": expense.accountant_comment, "reviewed_by_name": expense.reviewed_by.full_name if expense.reviewed_by else None,
        "created_at": expense.created_at.isoformat(), "updated_at": expense.updated_at.isoformat(),
    }
    if detailed:
        data["attachments"] = [{"id": a.id, "document_type": a.document_type, "original_name": a.original_name, "file_size": a.file_size, "created_at": a.created_at.isoformat()} for a in expense.attachments.all()]
        data["payments"] = [{"id": p.id, "amount": str(p.amount), "paid_at": p.paid_at.isoformat(), "payment_method": p.payment_method, "transaction_number": p.transaction_number} for p in expense.payments.all()]
        data["history"] = [{"id": x.id, "action": x.action, "old_status": x.old_status, "new_status": x.new_status, "comment": x.comment, "user_name": x.user.full_name, "created_at": x.created_at.isoformat()} for x in expense.activity_logs.all()]
    return data


def _get_expense(expense_id, lock=False):
    try:
        query = ExpenseModel.objects.select_related("department", "created_by", "reviewed_by").prefetch_related("attachments", "payments", "activity_logs__user")
        if lock: query = query.select_for_update(of=("self",))
        return query.get(id=expense_id)
    except ExpenseModel.DoesNotExist:
        raise EntityNotFoundException("Expense", expense_id)


def _log(expense, user, action, old_status="", new_status="", comment=""):
    ExpenseActivityLogModel.objects.create(expense=expense, user_id=user.id, action=action, old_status=old_status, new_status=new_status, comment=comment)


def _notify(user_id, title, message, expense_id):
    if user_id:
        NotificationModel.objects.create(user_id=user_id, title=title, message=message, entity_type="expense", entity_id=expense_id)


def _notify_accountants(expense, title, message):
    recipients = UserModel.objects.filter(role__in=["accountant", "general_director"], is_active=True).exclude(id=expense.created_by_id)
    for recipient in recipients:
        _notify(recipient.id, title, message, expense.id)


def _apply_expense_filters(queryset, params, allow_department=True):
    """Validate query parameters before applying ORM filters to avoid accidental 500 responses."""
    for key in ("start_date", "end_date"):
        if params.get(key):
            try:
                date.fromisoformat(params[key])
            except ValueError:
                raise ValidationException(f"{key} must be in YYYY-MM-DD format")
            queryset = queryset.filter(**{f"purchase_date__{'gte' if key == 'start_date' else 'lte'}": params[key]})
    if params.get("department_id") and allow_department:
        try:
            department_id = int(params["department_id"])
        except (TypeError, ValueError):
            raise ValidationException("department_id must be an integer")
        queryset = queryset.filter(department_id=department_id)
    return queryset


def _expense_department_id(user, requested_department_id):
    """Resolve an expense department without allowing staff to impersonate another unit."""
    if str(user.role) not in FINANCE_ROLES:
        if not user.department_id:
            raise ValidationException("Your account must be assigned to a department before creating an expense")
        return user.department_id

    if requested_department_id in (None, ""):
        return user.department_id
    try:
        department_id = int(requested_department_id)
    except (TypeError, ValueError):
        raise ValidationException("department_id must be an integer")
    if department_id <= 0 or not DepartmentModel.objects.filter(id=department_id).exists():
        raise ValidationException("department_id must reference an existing department")
    return department_id


@csrf_exempt
@login_required_domain
@finance_transaction
def expenses_view(request):
    if request.method == "GET":
        qs = ExpenseModel.objects.select_related("department", "created_by", "reviewed_by")
        user = request.user_entity
        if not _can_manage(user):
            if str(user.role) == "manager" and user.department_id:
                qs = qs.filter(department_id=user.department_id)
            else:
                qs = qs.filter(created_by_id=user.id)
        for field in ("status", "category", "payment_source", "currency"):
            if request.GET.get(field): qs = qs.filter(**{field: request.GET[field]})
        qs = _apply_expense_filters(qs, request.GET, allow_department=_can_manage(user))
        return JsonResponse({"expenses": [_serialize(x) for x in qs.order_by("-purchase_date", "-id")]})
    if request.method != "POST": return JsonResponse({"error": "Method not allowed"}, status=405)
    try: body = json.loads(request.body or "{}")
    except json.JSONDecodeError: raise ValidationException("Invalid JSON body")
    required = ["purchase_date", "category", "amount", "payment_source", "payment_method"]
    if any(not body.get(key) for key in required): raise ValidationException("purchase_date, category, amount, payment_source and payment_method are required")
    try:
        amount = Decimal(str(body["amount"])); purchase_date = date.fromisoformat(body["purchase_date"])
    except (InvalidOperation, ValueError): raise ValidationException("Invalid amount or purchase_date")
    if amount <= 0 or body["payment_source"] not in {"company", "personal"}: raise ValidationException("Amount must be positive and payment_source valid")
    department_id = _expense_department_id(request.user_entity, body.get("department_id"))
    require_open_date(purchase_date)
    expense = ExpenseModel.objects.create(purchase_date=purchase_date, department_id=department_id, category=str(body["category"]).strip(), supplier=str(body.get("supplier", "")).strip(), description=str(body.get("description", "")).strip(), amount=amount, currency=str(body.get("currency", "KGS")).upper(), payment_source=body["payment_source"], payment_method=body["payment_method"], status="draft", created_by_id=request.user_entity.id)
    _log(expense, request.user_entity, "created", new_status="draft")
    return JsonResponse(_serialize(_get_expense(expense.id), True), status=201)


@csrf_exempt
@login_required_domain
@finance_transaction
def expense_detail_view(request, expense_id):
    expense = _get_expense(expense_id, lock=request.method == "PUT")
    if not _can_access(request.user_entity, expense): raise PermissionDeniedException("Access to this expense is denied")
    if request.method == "GET": return JsonResponse(_serialize(expense, True))
    if request.method == "PUT":
        require_open_date(expense.purchase_date)
        if expense.created_by_id != request.user_entity.id or expense.status not in EDITABLE_STATUSES: raise PermissionDeniedException("Only your draft or returned expense can be edited")
        try: body = json.loads(request.body or "{}")
        except json.JSONDecodeError: raise ValidationException("Invalid JSON body")
        for field in ("category", "supplier", "description", "currency", "payment_method"):
            if field in body: setattr(expense, field, str(body[field]).strip())
        if "amount" in body:
            try: expense.amount = Decimal(str(body["amount"]))
            except InvalidOperation: raise ValidationException("Invalid amount")
            if expense.amount <= 0: raise ValidationException("Amount must be positive")
        expense.save(); _log(expense, request.user_entity, "updated")
        return JsonResponse(_serialize(_get_expense(expense.id), True))
    return JsonResponse({"error": "Method not allowed"}, status=405)


@csrf_exempt
@login_required_domain
@finance_transaction
def expense_status_view(request, expense_id):
    if request.method != "POST": return JsonResponse({"error": "Method not allowed"}, status=405)
    expense = _get_expense(expense_id, lock=True); user = request.user_entity
    try: body = json.loads(request.body or "{}")
    except json.JSONDecodeError: raise ValidationException("Invalid JSON body")
    status = body.get("status"); comment = str(body.get("comment", "")).strip()
    if status not in VALID_STATUSES: raise ValidationException("Invalid status")
    if status == "under_review":
        if expense.created_by_id != user.id or expense.status not in EDITABLE_STATUSES: raise PermissionDeniedException("Only your draft or returned expense can be submitted")
    else:
        if not _can_manage(user): raise PermissionDeniedException("Only accountant or director can change this status")
        if status == "rejected" and not comment: raise ValidationException("Comment is required when rejecting an expense")
    if status not in ALLOWED_STATUS_TRANSITIONS.get(expense.status, set()):
        raise ValidationException(f"Status cannot change from {expense.status} to {status}")
    require_open_date(expense.purchase_date)
    old_status = expense.status; expense.status = status; expense.accountant_comment = comment if _can_manage(user) else expense.accountant_comment
    if _can_manage(user): expense.reviewed_by_id = user.id; expense.reviewed_at = timezone.now()
    expense.save(); _log(expense, user, "status_changed", old_status, status, comment)
    if status == "under_review":
        _notify_accountants(expense, "Новая заявка на расход", f"{expense.created_by.full_name} отправил(а) расход «{expense.category}» на сумму {expense.amount} {expense.currency}.")
    elif _can_manage(user):
        _notify(expense.created_by_id, "Статус расхода изменён", f"Заявка «{expense.category}» — {dict(ExpenseModel.STATUS_CHOICES)[status]}. {comment}".strip(), expense.id)
    return JsonResponse(_serialize(_get_expense(expense.id), True))


@csrf_exempt
@login_required_domain
@finance_transaction
def expense_attachments_view(request, expense_id):
    if request.method != "POST": return JsonResponse({"error": "Method not allowed"}, status=405)
    expense = _get_expense(expense_id, lock=True)
    if expense.created_by_id != request.user_entity.id and not _can_manage(request.user_entity): raise PermissionDeniedException("Cannot attach documents to this expense")
    if expense.status in TERMINAL_ATTACHMENT_STATUSES:
        raise ValidationException("Documents cannot be added after the expense has a final decision")
    require_open_date(expense.purchase_date)
    upload = request.FILES.get("file")
    if not upload: raise ValidationException("file is required")
    if upload.size > 15 * 1024 * 1024: raise ValidationException("File size must not exceed 15 MB")
    allowed = {"application/pdf", "image/jpeg", "image/png", "image/webp"}
    if upload.content_type not in allowed: raise ValidationException("Only PDF, JPG, PNG and WEBP documents are allowed")
    document_type = request.POST.get("document_type", "other")
    if document_type not in {"kkm_receipt", "bank_receipt", "invoice", "waybill", "other"}: raise ValidationException("Invalid document_type")
    item = ExpenseAttachmentModel.objects.create(expense=expense, document_type=document_type, file=upload, original_name=upload.name[:255], file_size=upload.size, uploaded_by_id=request.user_entity.id)
    _log(expense, request.user_entity, "document_attached", comment=item.original_name)
    return JsonResponse({"id": item.id, "document_type": item.document_type, "original_name": item.original_name, "file_size": item.file_size}, status=201)


@login_required_domain
def expense_attachment_download_view(request, attachment_id):
    try: attachment = ExpenseAttachmentModel.objects.select_related("expense__department", "expense__created_by").get(id=attachment_id)
    except ExpenseAttachmentModel.DoesNotExist: raise EntityNotFoundException("Expense attachment", attachment_id)
    if not _can_access(request.user_entity, attachment.expense): raise PermissionDeniedException("Access to this document is denied")
    return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=attachment.original_name)


@csrf_exempt
@login_required_domain
@finance_transaction
def reimbursement_payment_view(request, expense_id):
    if request.method != "POST": return JsonResponse({"error": "Method not allowed"}, status=405)
    if not _can_manage(request.user_entity): raise PermissionDeniedException("Only accountant or director can record payment")
    try: body = json.loads(request.body or "{}")
    except json.JSONDecodeError: raise ValidationException("Invalid JSON body")
    with transaction.atomic():
        try:
            expense = ExpenseModel.objects.select_for_update(of=("self",)).select_related("department", "created_by", "reviewed_by").get(id=expense_id)
        except ExpenseModel.DoesNotExist:
            raise EntityNotFoundException("Expense", expense_id)
        if expense.payment_source != "personal" or expense.status != "approved": raise ValidationException("Only approved personal expenses can be paid")
        try: amount = Decimal(str(body.get("amount", expense.amount)))
        except InvalidOperation: raise ValidationException("Invalid amount")
        paid_total = ReimbursementPaymentModel.objects.filter(expense_id=expense.id).aggregate(total=Sum("amount"))["total"] or Decimal("0")
        if amount <= 0 or paid_total + amount > expense.amount: raise ValidationException("Payment amount exceeds reimbursement balance")
        paid_at = timezone.now()
        require_open_date(business_date(paid_at))
        payment = ReimbursementPaymentModel.objects.create(expense=expense, paid_at=paid_at, amount=amount, payment_method=str(body.get("payment_method", "bank_transfer")), transaction_number=str(body.get("transaction_number", "")), created_by_id=request.user_entity.id)
        if paid_total + amount == expense.amount:
            old = expense.status; expense.status = "paid"; expense.save(); _log(expense, request.user_entity, "payment_recorded", old, "paid", payment.transaction_number)
        else: _log(expense, request.user_entity, "partial_payment", comment=payment.transaction_number)
        _notify(expense.created_by_id, "Возмещение выплачено", f"По расходу «{expense.category}» зафиксирована выплата {amount} {expense.currency}.", expense.id)
    return JsonResponse(_serialize(_get_expense(expense.id), True), status=201)


@login_required_domain
def expense_analytics_view(request):
    if not _can_manage(request.user_entity): raise PermissionDeniedException("Finance analytics is available to accountant and director only")
    qs = ExpenseModel.objects.all()
    qs = _apply_expense_filters(qs, request.GET)
    summary = qs.aggregate(total=Sum("amount"), count=Count("id"), personal=Sum("amount", filter=Q(payment_source="personal")), pending=Sum("amount", filter=Q(payment_source="personal", status__in=["under_review", "needs_info", "approved"])))
    def grouped(field): return list(qs.values(field).annotate(amount=Sum("amount"), count=Count("id")).order_by("-amount"))
    return JsonResponse({"summary": {k: str(v or 0) for k, v in summary.items()}, "by_category": grouped("category"), "by_department": list(qs.values("department__name").annotate(amount=Sum("amount"), count=Count("id")).order_by("-amount")), "by_employee": list(qs.values("created_by__full_name").annotate(amount=Sum("amount"), count=Count("id")).order_by("-amount"))})


@login_required_domain
def expense_export_view(request):
    if not _can_manage(request.user_entity): raise PermissionDeniedException("Finance reports are available to accountant and director only")
    fmt = request.GET.get("format", "xlsx")
    if fmt not in {"xlsx", "docx", "pdf"}: raise ValidationException("format must be xlsx, docx or pdf")
    qs = ExpenseModel.objects.select_related("department", "created_by").all()
    qs = _apply_expense_filters(qs, request.GET)
    rows = [[x.purchase_date.isoformat(), x.department.name if x.department else "—", x.created_by.full_name, x.category, x.supplier, str(x.amount), x.currency, x.payment_source, dict(ExpenseModel.STATUS_CHOICES)[x.status]] for x in qs]
    headers = ["Дата", "Отдел", "Сотрудник", "Категория", "Поставщик", "Сумма", "Валюта", "Источник", "Статус"]
    if fmt == "xlsx":
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Расходы"; ws.append(headers)
        for cell in ws[1]: cell.font = Font(bold=True)
        for row in rows: ws.append(row)
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 2, 32)
        output = io.BytesIO(); wb.save(output); content, content_type = output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "docx":
        from docx import Document
        document = Document(); document.add_heading("Отчёт по расходам", 0); table = document.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, headers): cell.text = value
        for row in rows:
            for cell, value in zip(table.add_row().cells, row): cell.text = value
        output = io.BytesIO(); document.save(output); content, content_type = output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
        from reportlab.lib import colors
        output = io.BytesIO(); table = Table([headers] + rows, repeatRows=1); table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .25, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("FONTSIZE", (0, 0), (-1, -1), 7)])); SimpleDocTemplate(output, pagesize=landscape(A4)).build([table]); content, content_type = output.getvalue(), "application/pdf"
    response = HttpResponse(content, content_type=content_type); response["Content-Disposition"] = f'attachment; filename="expenses.{fmt}"'; return response
