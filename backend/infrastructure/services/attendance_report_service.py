import io
import calendar
from datetime import datetime, date, time, timedelta
from typing import Dict, Any, List, Optional, Tuple
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from backend.infrastructure.persistence.models import (
    UserModel, DepartmentModel, BoardModel, BoardMemberModel,
    AttendanceRecordModel, AbsenceModel, WorkScheduleModel, OfficeModel
)


class AttendanceReportService:
    def get_monthly_timesheet_data(
        self,
        year: int,
        month: int,
        department_id: Optional[int] = None,
        user_id: Optional[int] = None,
        current_user: Optional[Any] = None
    ) -> Dict[str, Any]:
        """Calculates detailed monthly timesheet and statistics for all employees or filtered subset"""
        num_days = calendar.monthrange(year, month)[1]
        start_date = timezone.make_aware(datetime(year, month, 1, 0, 0, 0))
        end_date = timezone.make_aware(datetime(year, month, num_days, 23, 59, 59))

        user_qs = UserModel.objects.select_related("department").filter(is_active=True)
        if user_id:
            user_qs = user_qs.filter(id=user_id)
        elif department_id:
            user_qs = user_qs.filter(department_id=department_id)

        # RBAC check: if manager, restrict to board members
        if current_user and current_user.is_manager:
            board_ids = BoardMemberModel.objects.filter(user_id=current_user.id, role="manager").values_list("board_id", flat=True)
            member_user_ids = list(BoardMemberModel.objects.filter(board_id__in=board_ids).values_list("user_id", flat=True))
            if current_user.id not in member_user_ids:
                member_user_ids.append(current_user.id)
            user_qs = user_qs.filter(id__in=member_user_ids)
        elif current_user and current_user.is_employee:
            user_qs = user_qs.filter(id=current_user.id)

        users = list(user_qs.order_by("department__name", "full_name"))

        # Pre-fetch attendance records and absences for the entire month
        all_records = AttendanceRecordModel.objects.filter(
            timestamp__gte=start_date,
            timestamp__lte=end_date,
            verification_result__in=["verified", "manual_override"]
        ).order_by("timestamp")

        all_absences = AbsenceModel.objects.filter(
            start_time__lte=end_date
        ).filter(
            end_time__isnull=True
        ) | AbsenceModel.objects.filter(
            start_time__lte=end_date,
            end_time__gte=start_date
        )

        records_by_user_date: Dict[Tuple[int, int], List[AttendanceRecordModel]] = {}
        for r in all_records:
            k = (r.user_id, r.timestamp.day)
            records_by_user_date.setdefault(k, []).append(r)

        absences_by_user: Dict[int, List[AbsenceModel]] = {}
        for a in all_absences:
            absences_by_user.setdefault(a.user_id, []).append(a)

        # Pre-fetch schedules
        schedules_by_user = {s.user_id: s for s in WorkScheduleModel.objects.filter(user__isnull=False, is_active=True)}
        schedules_by_dept = {s.department_id: s for s in WorkScheduleModel.objects.filter(department__isnull=False, is_active=True)}
        default_schedule = WorkScheduleModel.objects.filter(user__isnull=True, department__isnull=True, is_active=True).first()

        employees_data = []
        department_summary: Dict[str, Dict[str, Any]] = {}

        for user in users:
            dept_name = user.department.name if user.department else "Без отдела"
            sched = schedules_by_user.get(user.id) or (schedules_by_dept.get(user.department_id) if user.department_id else None) or default_schedule
            sched_start = sched.start_time if sched else time(9, 0)
            sched_end = sched.end_time if sched else time(18, 0)

            total_worked_seconds = 0.0
            total_late_minutes = 0
            total_early_minutes = 0
            late_count = 0
            early_count = 0
            days_worked = 0

            break_lunch_seconds = 0.0
            break_business_seconds = 0.0
            break_personal_seconds = 0.0
            break_other_seconds = 0.0

            vacation_days = 0
            sick_days = 0
            trip_days = 0
            day_off_days = 0

            daily_records = []

            for day in range(1, num_days + 1):
                day_date = date(year, month, day)
                is_weekend = day_date.weekday() in (5, 6)
                day_records = records_by_user_date.get((user.id, day), [])

                # Check day status
                day_status = None
                user_absences = absences_by_user.get(user.id, [])
                for a in user_absences:
                    if a.category == "day_status":
                        a_start = timezone.localdate(a.start_time)
                        a_end = timezone.localdate(a.end_time) if a.end_time else a_start
                        if a_start <= day_date <= a_end:
                            day_status = a.absence_type
                            if a.absence_type == "vacation":
                                vacation_days += 1
                            elif a.absence_type == "sick_leave":
                                sick_days += 1
                            elif a.absence_type == "business_trip":
                                trip_days += 1
                            elif a.absence_type == "day_off":
                                day_off_days += 1
                            break

                check_in_r = next((r for r in day_records if r.record_type == "check_in"), None)
                check_out_r = next((r for r in reversed(day_records) if r.record_type == "check_out"), None)

                day_worked_hours = 0.0
                day_late_min = 0
                day_early_min = 0

                if check_in_r:
                    days_worked += 1
                    cin_local = timezone.localtime(check_in_r.timestamp)
                    sched_start_dt = cin_local.replace(hour=sched_start.hour, minute=sched_start.minute, second=0)

                    if cin_local > sched_start_dt:
                        diff = int((cin_local - sched_start_dt).total_seconds() / 60)
                        if diff > 5:  # 5 min grace threshold
                            day_late_min = diff
                            late_count += 1
                            total_late_minutes += diff

                    # Calculate checkout / worked duration
                    if check_out_r:
                        cout_local = timezone.localtime(check_out_r.timestamp)
                        sched_end_dt = cout_local.replace(hour=sched_end.hour, minute=sched_end.minute, second=0)
                        if cout_local < sched_end_dt:
                            diff_early = int((sched_end_dt - cout_local).total_seconds() / 60)
                            if diff_early > 5:
                                day_early_min = diff_early
                                early_count += 1
                                total_early_minutes += diff_early

                        raw_sec = max(0.0, (cout_local - cin_local).total_seconds())

                        # Intra-day breaks on that day
                        for a in user_absences:
                            if a.category == "break" and a.start_time.date() == day_date:
                                a_end = a.end_time or cout_local
                                a_sec = max(0.0, (a_end - a.start_time).total_seconds())
                                if a.absence_type == "business":
                                    break_business_seconds += a_sec
                                    # Business leave does not deduct from worked hours
                                elif a.absence_type == "lunch":
                                    break_lunch_seconds += a_sec
                                    raw_sec = max(0.0, raw_sec - a_sec)
                                elif a.absence_type == "personal":
                                    break_personal_seconds += a_sec
                                    raw_sec = max(0.0, raw_sec - a_sec)
                                else:
                                    break_other_seconds += a_sec
                                    raw_sec = max(0.0, raw_sec - a_sec)

                        total_worked_seconds += raw_sec
                        day_worked_hours = round(raw_sec / 3600.0, 1)
                    else:
                        # Only checked in
                        day_worked_hours = 0.0

                daily_records.append({
                    "day": day,
                    "date": day_date.isoformat(),
                    "is_weekend": is_weekend,
                    "check_in": timezone.localtime(check_in_r.timestamp).strftime("%H:%M") if check_in_r else None,
                    "check_out": timezone.localtime(check_out_r.timestamp).strftime("%H:%M") if check_out_r else None,
                    "worked_hours": day_worked_hours,
                    "late_minutes": day_late_min,
                    "early_minutes": day_early_min,
                    "day_status": day_status,
                })

            total_worked_hours = round(total_worked_seconds / 3600.0, 1)

            emp_entry = {
                "user_id": user.id,
                "full_name": user.full_name,
                "email": user.email,
                "department_name": dept_name,
                "role": user.role,
                "days_worked": days_worked,
                "total_worked_hours": total_worked_hours,
                "late_count": late_count,
                "total_late_minutes": total_late_minutes,
                "early_count": early_count,
                "total_early_minutes": total_early_minutes,
                "lunch_hours": round(break_lunch_seconds / 3600.0, 1),
                "business_hours": round(break_business_seconds / 3600.0, 1),
                "personal_hours": round(break_personal_seconds / 3600.0, 1),
                "other_hours": round(break_other_seconds / 3600.0, 1),
                "vacation_days": vacation_days,
                "sick_days": sick_days,
                "trip_days": trip_days,
                "day_off_days": day_off_days,
                "daily_records": daily_records,
            }
            employees_data.append(emp_entry)

            # Accumulate department summary
            if dept_name not in department_summary:
                department_summary[dept_name] = {
                    "department_name": dept_name,
                    "employees_count": 0,
                    "total_worked_hours": 0.0,
                    "total_late_count": 0,
                    "total_late_minutes": 0,
                    "total_early_count": 0,
                }
            department_summary[dept_name]["employees_count"] += 1
            department_summary[dept_name]["total_worked_hours"] = round(department_summary[dept_name]["total_worked_hours"] + total_worked_hours, 1)
            department_summary[dept_name]["total_late_count"] += late_count
            department_summary[dept_name]["total_late_minutes"] += total_late_minutes
            department_summary[dept_name]["total_early_count"] += early_count

        return {
            "year": year,
            "month": month,
            "month_name": calendar.month_name[month],
            "num_days": num_days,
            "total_employees": len(employees_data),
            "employees": employees_data,
            "departments": list(department_summary.values()),
        }

    def export_excel(self, data: Dict[str, Any]) -> io.BytesIO:
        """Generates rich Excel (.xlsx) file with monthly timesheet and summaries"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Табель за месяц"

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="0F172A")
        sub_font = Font(name="Arial", size=10, italic=True, color="475569")
        bold_font = Font(name="Arial", size=9, bold=True)
        regular_font = Font(name="Arial", size=9)

        thin_side = Side(style="thin", color="CBD5E1")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # Title
        ws["A1"] = f"ТАБЕЛЬ УЧЁТА РАБОЧЕГО ВРЕМЕНИ ЗА {data['month']:02d}.{data['year']}"
        ws["A1"].font = title_font
        ws["A2"] = f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws["A2"].font = sub_font

        # Summary Header
        headers = [
            "ФИО сотрудника",
            "Отдел",
            "Отработано дней",
            "Отработано часов",
            "Опозданий (раз)",
            "Опозданий (мин)",
            "Ранних уходов",
            "Обед (ч)",
            "По делам комп. (ч)",
            "Личное (ч)",
            "Отпуск (дн)",
            "Больничный (дн)",
            "Командировка (дн)",
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[row_idx].height = 28

        # Data Rows
        for emp in data["employees"]:
            row_idx += 1
            ws.cell(row=row_idx, column=1, value=emp["full_name"]).font = bold_font
            ws.cell(row=row_idx, column=2, value=emp["department_name"]).font = regular_font
            ws.cell(row=row_idx, column=3, value=emp["days_worked"]).font = regular_font
            ws.cell(row=row_idx, column=4, value=emp["total_worked_hours"]).font = bold_font
            ws.cell(row=row_idx, column=5, value=emp["late_count"]).font = regular_font
            ws.cell(row=row_idx, column=6, value=emp["total_late_minutes"]).font = regular_font
            ws.cell(row=row_idx, column=7, value=emp["early_count"]).font = regular_font
            ws.cell(row=row_idx, column=8, value=emp["lunch_hours"]).font = regular_font
            ws.cell(row=row_idx, column=9, value=emp["business_hours"]).font = regular_font
            ws.cell(row=row_idx, column=10, value=emp["personal_hours"]).font = regular_font
            ws.cell(row=row_idx, column=11, value=emp["vacation_days"]).font = regular_font
            ws.cell(row=row_idx, column=12, value=emp["sick_days"]).font = regular_font
            ws.cell(row=row_idx, column=13, value=emp["trip_days"]).font = regular_font

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=c)
                cell.border = border
                if c >= 3:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Output to buffer
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_word(self, data: Dict[str, Any]) -> io.BytesIO:
        """Generates Word (.docx) file with monthly timesheet"""
        doc = Document()
        doc.add_heading(f"Табель учёта рабочего времени за {data['month']:02d}.{data['year']}", level=1)
        doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Сотрудник", "Отдел", "Отработано дней", "Отработано (ч)", "Опоздания (раз/мин)", "Ранние уходы", "Отлучки (личн/дел)"]
        for i, t in enumerate(hdr_titles):
            hdr_cells[i].text = t
            p = hdr_cells[i].paragraphs[0]
            if p.runs:
                p.runs[0].font.bold = True

        for emp in data["employees"]:
            row_cells = table.add_row().cells
            row_cells[0].text = emp["full_name"]
            row_cells[1].text = emp["department_name"]
            row_cells[2].text = str(emp["days_worked"])
            row_cells[3].text = str(emp["total_worked_hours"])
            row_cells[4].text = f"{emp['late_count']} ({emp['total_late_minutes']}м)"
            row_cells[5].text = str(emp["early_count"])
            row_cells[6].text = f"{emp['personal_hours']}ч / {emp['business_hours']}ч"

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def export_pdf(self, data: Dict[str, Any]) -> io.BytesIO:
        """Generates PDF file in landscape mode with timesheet"""
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=10
        )

        sub_style = ParagraphStyle(
            name="SubStyle",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#475569"),
            spaceAfter=15
        )

        elements = [
            Paragraph(f"Табель учёта рабочего времени — {data['month']:02d}.{data['year']}", title_style),
            Paragraph(f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')} | Всего сотрудников: {data['total_employees']}", sub_style),
        ]

        table_data = [
            ["Сотрудник", "Отдел", "Дней", "Часов", "Опоздания", "Ранние", "Обед", "Дела", "Личное", "Отпуск/Больн."]
        ]

        for emp in data["employees"]:
            table_data.append([
                emp["full_name"][:20],
                emp["department_name"][:16],
                str(emp["days_worked"]),
                str(emp["total_worked_hours"]),
                f"{emp['late_count']} ({emp['total_late_minutes']}м)",
                str(emp["early_count"]),
                f"{emp['lunch_hours']}ч",
                f"{emp['business_hours']}ч",
                f"{emp['personal_hours']}ч",
                f"{emp['vacation_days']} / {emp['sick_days']}"
            ])

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 7.5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))

        elements.append(t)
        doc.build(elements)
        buf.seek(0)
        return buf
