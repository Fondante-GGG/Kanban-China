import io
from datetime import datetime, timedelta, timezone
from typing import Dict, Any, Optional, List
from django.utils import timezone as dj_timezone
from backend.application.interfaces.services import IReportService
from backend.infrastructure.persistence.models import (
    CardModel, ColumnModel, BoardModel, UserModel, DepartmentModel, BoardMemberModel
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


class ReportService(IReportService):
    def get_analytics_data(
        self,
        uow: Any,
        period: str = "month",
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        department_id: Optional[int] = None
    ) -> Dict[str, Any]:
        now = dj_timezone.now()

        # Determine date range
        if start_date and end_date:
            try:
                start_dt = datetime.fromisoformat(start_date.replace("Z", "+00:00"))
                end_dt = datetime.fromisoformat(end_date.replace("Z", "+00:00"))
                if start_dt.tzinfo is None:
                    start_dt = start_dt.replace(tzinfo=timezone.utc)
                if end_dt.tzinfo is None:
                    end_dt = end_dt.replace(tzinfo=timezone.utc)
            except Exception:
                start_dt = now - timedelta(days=30)
                end_dt = now
        elif period == "week":
            start_dt = now - timedelta(days=7)
            end_dt = now
        elif period == "quarter":
            start_dt = now - timedelta(days=90)
            end_dt = now
        else:  # month default
            start_dt = now - timedelta(days=30)
            end_dt = now

        # Base card queryset
        card_qs = CardModel.objects.select_related("column__board__department").prefetch_related("assignee_relations__user")
        # Reports must respect the selected reporting range.  Filtering here also
        # prevents unrelated historical cards from changing current-period KPIs.
        card_qs = card_qs.filter(created_at__gte=start_dt, created_at__lte=end_dt)
        if department_id:
            card_qs = card_qs.filter(column__board__department_id=department_id)

        all_cards = list(card_qs)

        total_cards = len(all_cards)
        completed_cards = 0
        overdue_cards = 0
        in_progress_cards = 0
        bugs_cards = 0

        # Department metrics map
        dept_map: Dict[int, Dict[str, Any]] = {}
        for dept in DepartmentModel.objects.all():
            dept_map[dept.id] = {
                "id": dept.id,
                "name": dept.name,
                "total": 0,
                "completed": 0,
                "overdue": 0,
                "in_progress": 0,
            }

        # Employee metrics map
        emp_map: Dict[int, Dict[str, Any]] = {}
        for user in UserModel.objects.select_related("department").all():
            emp_map[user.id] = {
                "id": user.id,
                "full_name": user.full_name,
                "email": user.email,
                "role": user.role,
                "department_name": user.department.name if user.department else "—",
                "assigned_count": 0,
                "completed_count": 0,
                "overdue_count": 0,
                "in_progress_count": 0,
                "completion_times_hours": [],
            }

        for card in all_cards:
            col_name = (card.column.title_ru or "").lower()
            is_completed = "готов" in col_name or "done" in col_name
            is_bug = "баг" in col_name or "bug" in col_name
            is_in_progress = "процесс" in col_name or "progress" in col_name or "проверк" in col_name

            is_overdue = False
            if card.due_date and card.due_date < now and not is_completed:
                is_overdue = True

            if is_completed:
                completed_cards += 1
            if is_bug:
                bugs_cards += 1
            if is_in_progress:
                in_progress_cards += 1
            if is_overdue:
                overdue_cards += 1

            # Dept aggregation
            card_dept_id = card.column.board.department_id if card.column.board else None
            if card_dept_id and card_dept_id in dept_map:
                dept_map[card_dept_id]["total"] += 1
                if is_completed:
                    dept_map[card_dept_id]["completed"] += 1
                if is_overdue:
                    dept_map[card_dept_id]["overdue"] += 1
                if is_in_progress:
                    dept_map[card_dept_id]["in_progress"] += 1

            # Employee aggregation
            assignee_relations = card.assignee_relations.all()
            for rel in assignee_relations:
                u_id = rel.user_id
                if u_id in emp_map:
                    emp_map[u_id]["assigned_count"] += 1
                    if is_completed:
                        emp_map[u_id]["completed_count"] += 1
                        diff_hours = (card.updated_at - card.created_at).total_seconds() / 3600.0
                        emp_map[u_id]["completion_times_hours"].append(diff_hours)
                    if is_overdue:
                        emp_map[u_id]["overdue_count"] += 1
                    if is_in_progress:
                        emp_map[u_id]["in_progress_count"] += 1

        # Format employee list
        emp_list = []
        for emp in emp_map.values():
            times = emp["completion_times_hours"]
            avg_hours = round(sum(times) / len(times), 1) if times else 0
            rate = round((emp["completed_count"] / emp["assigned_count"] * 100), 1) if emp["assigned_count"] > 0 else 0
            emp_list.append({
                "id": emp["id"],
                "full_name": emp["full_name"],
                "email": emp["email"],
                "role": emp["role"],
                "department_name": emp["department_name"],
                "assigned_count": emp["assigned_count"],
                "completed_count": emp["completed_count"],
                "overdue_count": emp["overdue_count"],
                "in_progress_count": emp["in_progress_count"],
                "avg_completion_hours": avg_hours,
                "completion_rate": rate,
            })

        emp_list.sort(key=lambda x: x["assigned_count"], reverse=True)

        dept_list = list(dept_map.values())
        for d in dept_list:
            d["completion_rate"] = round((d["completed"] / d["total"] * 100), 1) if d["total"] > 0 else 0
        dept_list.sort(key=lambda x: x["total"], reverse=True)

        overall_rate = round((completed_cards / total_cards * 100), 1) if total_cards > 0 else 0

        return {
            "period": period,
            "start_date": start_dt.strftime("%Y-%m-%d"),
            "end_date": end_dt.strftime("%Y-%m-%d"),
            "summary": {
                "total_tasks": total_cards,
                "completed_tasks": completed_cards,
                "in_progress_tasks": in_progress_cards,
                "overdue_tasks": overdue_cards,
                "bug_tasks": bugs_cards,
                "completion_rate": overall_rate,
            },
            "departments": dept_list,
            "employees": emp_list,
        }

    def export_excel(self, data: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        subhead_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        subhead_font = Font(name="Arial", size=11, bold=True, color="1E3A8A")
        bold_font = Font(name="Arial", size=10, bold=True)
        regular_font = Font(name="Arial", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        # Sheet 1: Summary
        ws_sum = wb.active
        ws_sum.title = "Сводка"
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum["A1"] = "ОТЧЁТ ПО ЭФФЕКТИВНОСТИ И ЗАДАЧАМ КОМПАНИИ"
        ws_sum["A1"].font = Font(name="Arial", size=16, bold=True, color="1E3A8A")
        ws_sum["A2"] = f"Период: {data['start_date']} — {data['end_date']} ({data['period']})"
        ws_sum["A2"].font = Font(name="Arial", size=11, italic=True)

        summary_metrics = [
            ("Всего задач", data["summary"]["total_tasks"]),
            ("Выполнено задач", data["summary"]["completed_tasks"]),
            ("В процессе", data["summary"]["in_progress_tasks"]),
            ("Просрочено задач", data["summary"]["overdue_tasks"]),
            ("Задач типа Баги", data["summary"]["bug_tasks"]),
            ("Процент выполнения (%)", f"{data['summary']['completion_rate']}%"),
        ]

        ws_sum["A4"] = "Метрика"
        ws_sum["B4"] = "Значение"
        ws_sum["A4"].fill = header_fill
        ws_sum["A4"].font = header_font
        ws_sum["B4"].fill = header_fill
        ws_sum["B4"].font = header_font

        row = 5
        for metric, val in summary_metrics:
            ws_sum.cell(row=row, column=1, value=metric).font = regular_font
            ws_sum.cell(row=row, column=1).border = thin_border
            cell_val = ws_sum.cell(row=row, column=2, value=val)
            cell_val.font = bold_font
            cell_val.alignment = center_align
            cell_val.border = thin_border
            row += 1

        # Sheet 2: Departments
        ws_dept = wb.create_sheet(title="Отделы")
        ws_dept.views.sheetView[0].showGridLines = True
        dept_headers = ["Отдел", "Всего задач", "Выполнено", "В процессе", "Просрочено", "Успешность (%)"]
        for col_idx, h in enumerate(dept_headers, start=1):
            cell = ws_dept.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for r_idx, d in enumerate(data["departments"], start=2):
            ws_dept.cell(row=r_idx, column=1, value=d["name"]).alignment = left_align
            ws_dept.cell(row=r_idx, column=2, value=d["total"]).alignment = center_align
            ws_dept.cell(row=r_idx, column=3, value=d["completed"]).alignment = center_align
            ws_dept.cell(row=r_idx, column=4, value=d["in_progress"]).alignment = center_align
            ws_dept.cell(row=r_idx, column=5, value=d["overdue"]).alignment = center_align
            ws_dept.cell(row=r_idx, column=6, value=f"{d['completion_rate']}%").alignment = center_align
            for c_idx in range(1, 7):
                ws_dept.cell(row=r_idx, column=c_idx).border = thin_border
                ws_dept.cell(row=r_idx, column=c_idx).font = regular_font

        # Sheet 3: Employees
        ws_emp = wb.create_sheet(title="Сотрудники")
        ws_emp.views.sheetView[0].showGridLines = True
        emp_headers = ["ФИО", "Email", "Отдел", "Назначено", "Выполнено", "Просрочено", "В процессе", "Ср. время (ч)", "Успешность (%)"]
        for col_idx, h in enumerate(emp_headers, start=1):
            cell = ws_emp.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for r_idx, emp in enumerate(data["employees"], start=2):
            ws_emp.cell(row=r_idx, column=1, value=emp["full_name"]).alignment = left_align
            ws_emp.cell(row=r_idx, column=2, value=emp["email"]).alignment = left_align
            ws_emp.cell(row=r_idx, column=3, value=emp["department_name"]).alignment = left_align
            ws_emp.cell(row=r_idx, column=4, value=emp["assigned_count"]).alignment = center_align
            ws_emp.cell(row=r_idx, column=5, value=emp["completed_count"]).alignment = center_align
            ws_emp.cell(row=r_idx, column=6, value=emp["overdue_count"]).alignment = center_align
            ws_emp.cell(row=r_idx, column=7, value=emp["in_progress_count"]).alignment = center_align
            ws_emp.cell(row=r_idx, column=8, value=emp["avg_completion_hours"]).alignment = center_align
            ws_emp.cell(row=r_idx, column=9, value=f"{emp['completion_rate']}%").alignment = center_align
            for c_idx in range(1, 10):
                ws_emp.cell(row=r_idx, column=c_idx).border = thin_border
                ws_emp.cell(row=r_idx, column=c_idx).font = regular_font

        # Auto-adjust column widths
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out.getvalue()

    def export_word(self, data: Dict[str, Any]) -> bytes:
        doc = Document()

        # Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("Отчёт по задачам и производительности")
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 58, 138)

        p_meta = doc.add_paragraph()
        p_meta.add_run(f"Период: {data['start_date']} — {data['end_date']} ({data['period']})").italic = True

        doc.add_heading("1. Сводные показатели компании", level=1)
        sum_table = doc.add_table(rows=1, cols=2)
        sum_table.style = 'Light Shading Accent 1'
        hdr_cells = sum_table.rows[0].cells
        hdr_cells[0].text = "Показатель"
        hdr_cells[1].text = "Значение"

        metrics = [
            ("Всего задач", str(data["summary"]["total_tasks"])),
            ("Выполнено задач", str(data["summary"]["completed_tasks"])),
            ("В процессе", str(data["summary"]["in_progress_tasks"])),
            ("Просрочено", str(data["summary"]["overdue_tasks"])),
            ("Баги", str(data["summary"]["bug_tasks"])),
            ("Процент выполнения", f"{data['summary']['completion_rate']}%"),
        ]
        for name, val in metrics:
            row_cells = sum_table.add_row().cells
            row_cells[0].text = name
            row_cells[1].text = val

        doc.add_heading("2. Показатели по отделам", level=1)
        dept_table = doc.add_table(rows=1, cols=5)
        dept_table.style = 'Light Shading Accent 1'
        d_hdr = dept_table.rows[0].cells
        d_hdr[0].text = "Отдел"
        d_hdr[1].text = "Всего"
        d_hdr[2].text = "Выполнено"
        d_hdr[3].text = "Просрочено"
        d_hdr[4].text = "Успешность"

        for d in data["departments"]:
            row_cells = dept_table.add_row().cells
            row_cells[0].text = d["name"]
            row_cells[1].text = str(d["total"])
            row_cells[2].text = str(d["completed"])
            row_cells[3].text = str(d["overdue"])
            row_cells[4].text = f"{d['completion_rate']}%"

        doc.add_heading("3. Показатели по сотрудникам", level=1)
        emp_table = doc.add_table(rows=1, cols=6)
        emp_table.style = 'Light Shading Accent 1'
        e_hdr = emp_table.rows[0].cells
        e_hdr[0].text = "Сотрудник"
        e_hdr[1].text = "Отдел"
        e_hdr[2].text = "Задач"
        e_hdr[3].text = "Выполнено"
        e_hdr[4].text = "Ср. время (ч)"
        e_hdr[5].text = "Успешность"

        for emp in data["employees"]:
            row_cells = emp_table.add_row().cells
            row_cells[0].text = emp["full_name"]
            row_cells[1].text = emp["department_name"]
            row_cells[2].text = str(emp["assigned_count"])
            row_cells[3].text = str(emp["completed_count"])
            row_cells[4].text = str(emp["avg_completion_hours"])
            row_cells[5].text = f"{emp['completion_rate']}%"

        out = io.BytesIO()
        doc.save(out)
        out.seek(0)
        return out.getvalue()

    def export_pdf(self, data: Dict[str, Any]) -> bytes:
        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name='TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            name='SubTitleStyle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#4B5563'),
            spaceAfter=14
        )
        section_style = ParagraphStyle(
            name='SectionStyle',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#1E3A8A'),
            spaceBefore=12,
            spaceAfter=8
        )
        cell_style = ParagraphStyle(
            name='CellStyle',
            parent=styles['Normal'],
            fontSize=9,
            leading=11
        )
        header_cell_style = ParagraphStyle(
            name='HeaderCellStyle',
            parent=styles['Normal'],
            fontSize=9,
            leading=11,
            textColor=colors.white,
            fontName='Helvetica-Bold'
        )

        story = []
        story.append(Paragraph("Отчёт по задачам и производительности", title_style))
        story.append(Paragraph(f"Период: {data['start_date']} — {data['end_date']} ({data['period']})", subtitle_style))

        story.append(Paragraph("1. Сводные показатели", section_style))
        sum_rows = [
            [Paragraph("<b>Показатель</b>", header_cell_style), Paragraph("<b>Значение</b>", header_cell_style)],
            [Paragraph("Всего задач", cell_style), Paragraph(str(data["summary"]["total_tasks"]), cell_style)],
            [Paragraph("Выполнено", cell_style), Paragraph(str(data["summary"]["completed_tasks"]), cell_style)],
            [Paragraph("В процессе", cell_style), Paragraph(str(data["summary"]["in_progress_tasks"]), cell_style)],
            [Paragraph("Просрочено", cell_style), Paragraph(str(data["summary"]["overdue_tasks"]), cell_style)],
            [Paragraph("Баги", cell_style), Paragraph(str(data["summary"]["bug_tasks"]), cell_style)],
            [Paragraph("Процент выполнения", cell_style), Paragraph(f"{data['summary']['completion_rate']}%", cell_style)],
        ]
        sum_table = Table(sum_rows, colWidths=[200, 100])
        sum_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(sum_table)

        story.append(Paragraph("2. Показатели по отделам", section_style))
        dept_rows = [[
            Paragraph("<b>Отдел</b>", header_cell_style),
            Paragraph("<b>Всего</b>", header_cell_style),
            Paragraph("<b>Выполнено</b>", header_cell_style),
            Paragraph("<b>Просрочено</b>", header_cell_style),
            Paragraph("<b>Успешность</b>", header_cell_style)
        ]]
        for d in data["departments"]:
            dept_rows.append([
                Paragraph(d["name"], cell_style),
                Paragraph(str(d["total"]), cell_style),
                Paragraph(str(d["completed"]), cell_style),
                Paragraph(str(d["overdue"]), cell_style),
                Paragraph(f"{d['completion_rate']}%", cell_style)
            ])
        dept_table = Table(dept_rows, colWidths=[180, 80, 80, 80, 100])
        dept_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(dept_table)

        story.append(Paragraph("3. Показатели по сотрудникам", section_style))
        emp_rows = [[
            Paragraph("<b>Сотрудник</b>", header_cell_style),
            Paragraph("<b>Отдел</b>", header_cell_style),
            Paragraph("<b>Задач</b>", header_cell_style),
            Paragraph("<b>Выполнено</b>", header_cell_style),
            Paragraph("<b>Ср. время</b>", header_cell_style),
            Paragraph("<b>Успешность</b>", header_cell_style)
        ]]
        for emp in data["employees"][:25]:  # top 25
            emp_rows.append([
                Paragraph(emp["full_name"], cell_style),
                Paragraph(emp["department_name"], cell_style),
                Paragraph(str(emp["assigned_count"]), cell_style),
                Paragraph(str(emp["completed_count"]), cell_style),
                Paragraph(f"{emp['avg_completion_hours']} ч", cell_style),
                Paragraph(f"{emp['completion_rate']}%", cell_style)
            ])
        emp_table = Table(emp_rows, colWidths=[130, 110, 60, 70, 70, 80])
        emp_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(emp_table)

        doc.build(story)
        out.seek(0)
        return out.getvalue()
